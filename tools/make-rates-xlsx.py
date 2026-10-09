# Makes deliver/Mangal-rates.xlsx from src/data/rates.csv: the sheet the owner keeps in
# Google Sheets. Yellow boxes are rates that could not be read clearly from the photos.
import csv
import os
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
rows = list(csv.reader(open(os.path.join(ROOT, "src/data/rates.csv"), encoding="utf-8")))

F = lambda **k: Font(name="Arial", size=11, **k)
YELLOW = PatternFill("solid", fgColor="FFF200")
HEAD = PatternFill("solid", fgColor="0F1E33")
SET = PatternFill("solid", fgColor="E6EBF2")
thin = Side(style="thin", color="B9C6D6")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

# (brand, type, item, size) -> what the owner needs to check
doubt = {
    ("Ashirvad", "Plain", "Over Step Bend", "3/4"): "Overwritten on your sheet. It looks like 107. Please type the correct rate.",
    ("Ashirvad", "Plain", "Over Step Bend", "1"): "Overwritten on your sheet. It looks like 265. Please type the correct rate.",
    ("Ashirvad", "Brass", "Elbow", "1 1/4"): "Not written on your sheet. Type the rate, or empty this box if you do not sell this size.",
    ("Ashirvad", "Brass", "Tee", "1 1/4"): "Not written on your sheet. Type the rate, or empty this box if you do not sell this size.",
    ("Ashirvad", "Brass", "MTA", "3/4x1/2"): "Your sheet has two rates here: H 283 and L 210. Which one should the app use? (If both are sold, tell me and I will add them as two items.)",
    ("Supreme", "Pipe", "Pipe SDR 11", "3/4"): 'It looks like 219.33, which seems low next to 658 for 1". Please type the correct rate.',
    ("Supreme", "Plain", "Socket", "1x3/4"): "Overwritten on your sheet. It looks like 36.90. Please type the correct rate.",
    ("Supreme", "Plain", "Ball Valve", "3/4"): "Smudged on your sheet. It looks like 220.10. Please type the correct rate.",
    ("Supreme", "Plain", "Union", "3/4"): "Empty on your sheet. Type the rate, or empty this box if you do not sell this size.",
    ("Supreme", "Plain", "Over Step Bend", "1 1/4"): 'Empty on your sheet. Type the rate, or empty this box if you do not sell this size. There is also a side note "3/4: 81.20, 1: 199.20" next to this row. Which item is that for?',
    ("Supreme", "Plain", "Elbow", "1 1/4x1"): "Empty on your sheet. Type the rate, or empty this box if you do not sell this size.",
    ("Supreme", "Brass", "Elbow", "1 1/4"): "It looks like 483.60 or 482.60. Please type the correct rate.",
    ("Supreme", "Brass", "MTA", "1 1/4"): "Overwritten on your sheet. It looks like 724.40. Please type the correct rate.",
    ("Supreme", "Brass", "FTA", "1x1/2"): "Overwritten on your sheet. It looks like 116.10. Please type the correct rate.",
}
for s in ("3/4", "1", "1x3/4", "3/4x1/2", "1x1/2"):
    doubt[("Ashirvad", "Brass", "Tee", s)] = "The Tee row on your sheet is overwritten and crossed out. Please type the correct rate."
for s in ("3/4", "1", "1 1/4"):
    doubt[("Supreme", "Plain", "MTA", s)] = 'Your MTA row is marked "PVC" with sizes 3/4x1/2 = 31.50, 1x3/4 = 33.60, 1 1/4x1 = 68. What are the plain MTA rates for 3/4, 1 and 1 1/4?'
    doubt[("Supreme", "Plain", "C Bend", s)] = 'These numbers are written next to "C Bend" under the Elbow row on your sheet. Please confirm they are the C Bend rates.'

wb = Workbook()
ws = wb.active
ws.title = "Rates"
sizes = None
used = set()
for r, cells in enumerate(rows, 1):
    if cells and cells[0].startswith("#"):
        for c, v in enumerate(cells, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = F(bold=(c == 1))
            cell.fill = SET
            cell.number_format = "@"
        continue
    if cells and cells[0] == "Brand":
        sizes = cells[3:]
        for c, v in enumerate(cells, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = F(bold=True, color="FFFFFF")
            cell.fill = HEAD
            cell.number_format = "@"
            cell.alignment = Alignment(horizontal="center" if c > 3 else "left")
            cell.border = box
        ws.freeze_panes = ws.cell(row=r + 1, column=4)
        continue
    brand, typ, item = cells[0], cells[1], cells[2]
    for c, v in enumerate(cells[:3], 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = F(bold=(c == 3))
        cell.border = box
    for i, size in enumerate(sizes):
        v = cells[3 + i] if 3 + i < len(cells) else ""
        cell = ws.cell(row=r, column=4 + i)
        cell.font = F()
        cell.border = box
        cell.alignment = Alignment(horizontal="right")
        if v == "?":
            cell.value = "?"
            cell.number_format = "@"
        else:
            cell.number_format = "0.00"
            if v != "":
                cell.value = float(v)
        key = (brand, typ, item, size)
        if key in doubt:
            used.add(key)
            cell.fill = YELLOW
            note = Comment(doubt[key], "Claude")
            note.width = 320
            note.height = 110
            cell.comment = note
        else:
            assert v != "?", f"? without a note: {key}"
assert used == set(doubt), set(doubt) - used
ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 10
ws.column_dimensions["C"].width = 20
for i in range(len(sizes)):
    ws.column_dimensions[get_column_letter(4 + i)].width = 11

how = wb.create_sheet("How to use")
lines = [
    ("How this sheet works", "title"),
    ('The app reads the "Rates" tab. One row is one item; each size has its own column.', None),
    ("An empty box means you do not sell that size. A box with ? means you sell it but the rate is not known yet: the app shows that size greyed out.", None),
    ("", None),
    ("Yellow boxes", "head"),
    ("A yellow box is a rate I could not read clearly from your photos, or want you to confirm. Hold the mouse over it (or tap it) to see what I read.", None),
    ("Type the correct rate in each yellow box, or empty the box if you do not sell that size.", None),
    ("", None),
    ("To change a rate", "head"),
    ('Type the new number in the box. Then change the date in the "#Rates updated" row at the top, so every phone shows the new date.', None),
    ("", None),
    ("To add or remove things", "head"),
    ("New item: add a row with Brand, Type and Item, then the rates. Example row:   Supreme | Plain | Reducer | 45.50 | 60 | 95", None),
    ("Remove an item: delete its row. Remove one size of an item: empty that box.", None),
    ("New size: type it as a new column heading after the last one, for example 1 1/2 or 2x1. Type an apostrophe first ('1 1/2) so it is not turned into a date.", None),
    ("New brand: just use the new brand name in column A. It appears in the app as a new brand button.", None),
    ("", None),
    ("Rules the app checks", "head"),
    ("Keep the heading row starting with Brand, Type, Item. Rates are numbers only (62 or 21.30). The same item and size must not appear twice.", None),
    ("Type decides the discount group: all Plain items of a brand share one discount, all Brass items another, Pipe another.", None),
    ('Items of type Brass are printed on the bill as "Tee (Brass)". Plain and Pipe items are printed without brackets.', None),
    ('If the sheet has a mistake, the app keeps using the last good rates and shows which row and column is wrong (tap "Rates updated" in the app).', None),
    ("", None),
    ("To connect this sheet to the app (one time)", "head"),
    ("1. Upload this file to Google Drive and open it with Google Sheets.", None),
    ('2. File > Share > Publish to web. Choose the "Rates" tab and "Comma-separated values (.csv)". Press Publish.', None),
    ("3. Copy the link it gives and send it to me. I will put it in the app. After that, every phone picks up your changes when it has internet (Google can take about 5 minutes).", None),
    ("", None),
    ("Source of the rates: your two rate-sheet photos dated 6 Oct 2026 (Ashirvad and Supreme). They are not confirmed yet.", None),
]
for i, (text, kind) in enumerate(lines, 1):
    c = how.cell(row=i, column=1, value=text)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.font = Font(name="Arial", size=16, bold=True) if kind == "title" else Font(name="Arial", size=12, bold=True) if kind == "head" else Font(name="Arial", size=11)
how.column_dimensions["A"].width = 120

q = wb.create_sheet("Please confirm")
q.append(["Brand", "Type", "Item", "Size", "What I need from you"])
for c in range(1, 6):
    cell = q.cell(row=1, column=c)
    cell.font = F(bold=True, color="FFFFFF")
    cell.fill = HEAD
order = {k: i for i, k in enumerate([(c[0], c[1], c[2]) for c in rows if c and not c[0].startswith("#") and c[0] != "Brand"])}
for key in sorted(doubt, key=lambda k: (order[(k[0], k[1], k[2])], sizes.index(k[3]))):
    q.append([key[0], key[1], key[2], key[3], doubt[key]])
q.append(["Ashirvad", "", "Reducing MTA", "", 'Your sheet has a hand-written note "Reducing MTA 54 ... 64" above the brass table. Which sizes are these, and should it be an item in the app?'])
for row in q.iter_rows(min_row=2):
    for cell in row:
        cell.font = F()
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.number_format = "@"
for col, w in zip("ABCDE", (12, 9, 18, 12, 110)):
    q.column_dimensions[col].width = w

os.makedirs(os.path.join(ROOT, "deliver"), exist_ok=True)
out = os.path.join(ROOT, "deliver/Mangal-rates.xlsx")
wb.save(out)
print("wrote", out, "| boxes to confirm:", len(doubt))
